"""Cross-check the preview renderer against the Power Fx generator's own geometry.

These are two independent implementations of the same layout, and until now nothing
compared them -- so the footer could sit 20px higher in Studio than in the preview and
every check still passed. Allan found that in the real app, which is precisely the failure
mode a preview is supposed to prevent.

Run after any change to spacing constants in either file.
"""
import re, sys
import openpyxl
from render_lib import compute_layout, PANELS

DATA = "OrgChartPosts_Data_v15.xlsx"
GEN = "gen_phase12_formula.py"

wb = openpyxl.load_workbook(DATA)
ws = wb.active
cols = [c.value for c in ws[1]]
posts = [dict(zip(cols, r)) for r in ws.iter_rows(min_row=2, values_only=True)]

panelY, content_h, base_viewH, rows = compute_layout(posts)
rowCY, rowCH = rows["rowCY"], rows["rowCH"]

src = open(GEN, encoding="utf-8").read()
def const(name):
    return int(re.search(name + r" = (\d+)", src).group(1))

top_pad, bot_pad, banner = const("FOOTER_TOP_PAD"), const("FOOTER_BOTTOM_PAD"), const("BANNER_H")

deepest = max(panelY[d] - 4 + 30 + content_h(d) + 16 for d in PANELS)
preview_bar_y = (base_viewH - 20) + 10
preview_total = (base_viewH - 20) + 10 + banner + 10
fx_bar_y = rowCY + 30 + rowCH + top_pad
fx_total = rowCY + 30 + rowCH + top_pad + banner + bot_pad

bad = 0
print(f"deepest panel box bottom : {deepest:.0f}")
for label, bar, total in (("preview", preview_bar_y, preview_total), ("power fx", fx_bar_y, fx_total)):
    clear = bar - deepest
    flag = "" if clear >= 8 else f"   <-- only {clear:.0f}px, footer overlaps or crowds the panels"
    if clear < 8:
        bad += 1
    print(f"  {label:<9} footer bar {bar:.0f}  total {total:.0f}  clearance {clear:+.0f}px{flag}")

if abs(preview_bar_y - fx_bar_y) > 0.5 or abs(preview_total - fx_total) > 0.5:
    print(f"  MISMATCH: preview and Power Fx disagree by {preview_bar_y - fx_bar_y:.0f}px "
          f"(bar) and {preview_total - fx_total:.0f}px (total height)")
    bad += 1

print("PREVIEW vs FORMULA:", "OK" if not bad else "PROBLEM")
sys.exit(1 if bad else 0)
